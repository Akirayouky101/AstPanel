#!/usr/bin/env python3
"""Converte il foglio "Contabilità ZG AST.xlsx" in INSERT SQL per contabilita_movimenti.

Ogni foglio è un mese (AGOSTO, SETTEMBRE...). Il blocco A-C è la prima azienda,
il blocco a destra (H-J o I-K) la seconda. Le sezioni CREDITORE = uscite,
DEBITORE/CLIENTE = entrate. Le voci con importo vuoto vengono saltate.

Uso:
    python3 import-contabilita-xlsx.py "/percorso/Contabilità ZG AST.xlsx" --sinistra ZG --destra AST > import-contabilita.sql
Poi eseguire il file generato nel SQL Editor di Supabase (dopo create-contabilita.sql).
"""
import argparse
import datetime as dt
import re
import sys

import openpyxl

MESI = {"GENNAIO": 1, "FEBBRAIO": 2, "MARZO": 3, "APRILE": 4, "MAGGIO": 5, "GIUGNO": 6, "LUGLIO": 7,
        "AGOSTO": 8, "SETTEMBRE": 9, "OTTOBRE": 10, "NOVEMBRE": 11, "DICEMBRE": 12}
TOTALI = ("TOT", "PREVISIONE", "USCITE EVASE", "EFFETTIVI", "PREV.UTILE", "SALDO", "RIMANZA")

CATEGORIE = [
    (r"STIPENDI|TFR", "stipendi"), (r"MULTA|SANZION", "multe"), (r"F24 IVA|IVA\b|TASS", "tasse"), (r"F24 CONTRIBUTI|CONTRIBUT|INPS|INAIL", "contributi"),
    (r"RATA|BANCA|BANK|LEASING|FINANZ", "finanziamento"), (r"ENILIVE|ENI\b|Q8|CARBUR", "carburante"),
]


def categoria(tipo, nome):
    if tipo == "entrata":
        return "fattura_cliente"
    for pattern, cat in CATEGORIE:
        if re.search(pattern, nome, re.I):
            return cat
    return "fornitore"


def sql_str(value):
    return "'" + str(value).replace("'", "''") + "'"


def azienda_sql(nome):
    return f"(SELECT id FROM public.aziende WHERE nome ILIKE {sql_str(nome)} OR ragione_sociale ILIKE {sql_str(nome + '%')} ORDER BY attiva DESC LIMIT 1)"


def leggi_blocco(ws, col_nome, col_importo, col_data, anno, mese):
    """Scorre una colonna: le intestazioni CREDITORE/DEBITORE/CLIENTE cambiano il tipo corrente."""
    tipo = None
    righe = []
    for r in range(1, ws.max_row + 1):
        nome = ws.cell(r, col_nome).value
        importo = ws.cell(r, col_importo).value
        data = ws.cell(r, col_data).value
        if isinstance(nome, str):
            testo = nome.strip().upper()
            if testo == "CREDITORE":
                tipo = "uscita"
                continue
            if testo in ("DEBITORE", "CLIENTE"):
                tipo = "entrata"
                continue
        if tipo is None or not isinstance(nome, str) or not nome.strip():
            continue
        if isinstance(importo, str) or importo is None or any(t in nome.upper() for t in TOTALI):
            continue
        if isinstance(data, dt.datetime):
            data = data.date()
        if not isinstance(data, dt.date):
            data = dt.date(anno, mese, 1)
        righe.append((tipo, nome.strip(), round(float(importo), 2), data))
    return righe


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("file")
    parser.add_argument("--sinistra", default="ZG", help="nome azienda del blocco A-C")
    parser.add_argument("--destra", default="AST", help="nome azienda del blocco a destra")
    parser.add_argument("--anno", type=int, default=None)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.file, data_only=True)
    out = ["BEGIN;"]
    totale = 0
    viste = set()  # le voci non incassate vengono ricopiate nel foglio successivo: si importano una volta sola
    for ws in wb.worksheets:
        mese = MESI.get(ws.title.strip().upper())
        if not mese or ws.max_row < 2:
            continue
        anno = args.anno
        if anno is None:
            date = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, dt.datetime)]
            anno = max(date).year if date else dt.date.today().year
        competenza = dt.date(anno, mese, 1)

        blocchi = [(args.sinistra, 1, 2, 3)]
        for col in range(5, ws.max_column + 1):
            intest = [str(ws.cell(r, col).value or "").strip().upper() for r in range(1, 4)]
            if "CREDITORE" in intest:
                blocchi.append((args.destra, col, col + 1, col + 2))
                break

        for azienda, cn, ci, cd in blocchi:
            righe = []
            for riga in leggi_blocco(ws, cn, ci, cd, anno, mese):
                chiave = (azienda, *riga)
                if chiave in viste:
                    continue
                viste.add(chiave)
                righe.append(riga)
            if not righe:
                continue
            out.append(f"\n-- {ws.title} {anno} · {azienda} ({len(righe)} voci)")
            for tipo, nome, importo, data in righe:
                out.append(
                    "INSERT INTO public.contabilita_movimenti (azienda_id, tipo, controparte, categoria, importo, data_riferimento, mese_competenza, stato, note) "
                    f"VALUES ({azienda_sql(azienda)}, {sql_str(tipo)}, {sql_str(nome)}, {sql_str(categoria(tipo, nome))}, {importo:.2f}, "
                    f"DATE '{data.isoformat()}', DATE '{competenza.isoformat()}', 'previsto', 'Import da Excel foglio {ws.title}');"
                )
                totale += 1
    out.append("\nCOMMIT;")
    out.append(f"-- {totale} movimenti generati. Verifica prima che le aziende '{args.sinistra}' e '{args.destra}' esistano in public.aziende.")
    print("\n".join(out))
    print(f"-- Nota: le voci importate sono tutte 'previsto'; segna pagate/incassate dall'app.", file=sys.stderr)


if __name__ == "__main__":
    main()
